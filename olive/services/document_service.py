from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
import logging
import tempfile

from ..models import DocumentRef
from ..utils.chunking import TextPage, chunk_pages
from ..utils.files import cache_file, content_hash, mime_type, stable_file_id
from .ocr_service import OCRService, is_ocr_candidate

logger = logging.getLogger(__name__)

try:
    import pdfplumber
except ImportError:  # pragma: no cover - optional import checked at runtime
    pdfplumber = None

try:
    from pypdf import PdfReader
except ImportError:  # pragma: no cover
    PdfReader = None

try:
    from docx import Document as DocxDocument
except ImportError:  # pragma: no cover
    DocxDocument = None


TEXT_SUFFIXES = {
    ".txt", ".md", ".csv", ".json", ".log", ".py", ".cs", ".java", ".js", ".ts",
    ".tsx", ".jsx", ".html", ".css", ".sql", ".xml", ".yaml", ".yml", ".toml", ".ini",
    ".ps1", ".bat", ".sh", ".cpp", ".c", ".h", ".hpp", ".go", ".rs",
}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}


@dataclass(slots=True)
class ExtractedDocument:
    ref: DocumentRef
    pages: list[TextPage]
    chunks: list[dict]


class DocumentService:
    def __init__(self, ocr: OCRService | None = None, cache_dir: Path | None = None):
        self.ocr = ocr or OCRService()
        self.cache_dir = cache_dir
    def supported_dialog_filters(self):
        return [
            ("Supported files", "*.pdf *.docx *.txt *.md *.csv *.json *.log *.py *.cs *.java *.js *.ts *.html *.css *.sql *.xml *.yaml *.yml *.png *.jpg *.jpeg *.gif *.webp *.bmp"),
            ("Documents", "*.pdf *.docx *.txt *.md *.csv *.json *.log"),
            ("Code", "*.py *.cs *.java *.js *.ts *.html *.css *.sql *.xml *.yaml *.yml"),
            ("Images", "*.png *.jpg *.jpeg *.gif *.webp *.bmp"),
            ("All Files", "*.*"),
        ]

    def is_image(self, path: Path) -> bool:
        return path.suffix.lower() in IMAGE_SUFFIXES

    def extract(self, path: Path, chat_id: str) -> ExtractedDocument:
        path = path.expanduser().resolve()
        if not path.exists() or not path.is_file():
            raise FileNotFoundError(path)

        file_id = f"{stable_file_id(path)}-{chat_id[:8]}"
        stored = cache_file(path, file_id, self.cache_dir)
        source_stat = path.stat()
        suffix = path.suffix.lower()

        if suffix == ".pdf":
            pages = self._extract_pdf(path)
            kind = "pdf"
        elif suffix == ".docx":
            pages = self._extract_docx(path)
            kind = "docx"
        elif suffix in TEXT_SUFFIXES or not suffix:
            pages = [TextPage(path.read_text(encoding="utf-8", errors="ignore"), None)]
            kind = "text"
        else:
            # Best-effort text read. Binary files will normally decode to little useful text,
            # so callers should show a clear failure if the extracted result is empty.
            raw = path.read_text(encoding="utf-8", errors="ignore")
            pages = [TextPage(raw, None)]
            kind = "text"

        cleaned = [TextPage(_normalise_page(p.text), p.page_number, p.origin_type, p.confidence)
                   for p in pages if p.text and p.text.strip()]
        if not cleaned and kind != "pdf":
            raise ValueError(f"No readable text was found in {path.name}")

        text_chunks = chunk_pages(cleaned)
        ref = DocumentRef(
            id=file_id,
            name=path.name,
            kind=kind,
            page_count=len(pages) if kind == "pdf" else None,
            unreadable_pages=[p.page_number for p in pages if not p.text.strip() and p.page_number],
            stored_path=str(stored),
            indexed=False,
            embedding_indexed=False,
            original_path=str(path), content_hash=content_hash(path),
            source_size=source_stat.st_size, source_mtime_ns=source_stat.st_mtime_ns,
        )
        chunks = [
            {
                "chat_id": chat_id,
                "document_name": path.name,
                "chunk_index": c.chunk_index,
                "page_number": c.page_number,
                "content": c.text,
                "origin_type": c.origin_type,
                "ocr_confidence": c.confidence,
            }
            for c in text_chunks
        ]
        return ExtractedDocument(ref=ref, pages=cleaned, chunks=chunks)

    def _extract_pdf(self, path: Path) -> list[TextPage]:
        pages: list[TextPage] = []
        if pdfplumber is not None:
            with pdfplumber.open(path) as pdf:
                for page_no, page in enumerate(pdf.pages, 1):
                    text = page.extract_text() or ""
                    tables = page.extract_tables() or []
                    origin = "native_pdf_text"
                    if tables:
                        table_text = []
                        for table in tables:
                            table_text.append(_table_to_markdown(table))
                        text = f"{text}\n\n" + "\n\n".join(t for t in table_text if t)
                        origin = "table_extraction"
                    confidence = None
                    image_count = len(getattr(page, "images", []) or [])
                    if is_ocr_candidate(text, image_count) and self.ocr.available:
                        try:
                            with tempfile.TemporaryDirectory() as tmp:
                                image_path = Path(tmp) / f"page-{page_no}.png"
                                page.to_image(resolution=200).save(str(image_path), format="PNG")
                                ocr_result = self.ocr.recognize(image_path)
                            if ocr_result.text.strip():
                                text, confidence, origin = ocr_result.text, ocr_result.confidence, "ocr"
                        except Exception:
                            logger.exception("OCR failed for PDF page %s in %s", page_no, path.name)
                    pages.append(TextPage(text=text, page_number=page_no,
                                          origin_type=origin, confidence=confidence))
            return _remove_repeated_page_margins(pages)

        if PdfReader is not None:
            reader = PdfReader(str(path))
            for page_no, page in enumerate(reader.pages, 1):
                pages.append(TextPage(text=page.extract_text() or "", page_number=page_no,
                                      origin_type="native_pdf_text"))
            return _remove_repeated_page_margins(pages)

        raise RuntimeError("PDF support is not installed. Run: pip install pdfplumber pypdf")

    def _extract_docx(self, path: Path) -> list[TextPage]:
        if DocxDocument is None:
            raise RuntimeError("DOCX support is not installed. Run: pip install python-docx")
        doc = DocxDocument(str(path))
        blocks: list[str] = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                blocks.append(paragraph.text.strip())
        for table in doc.tables:
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            blocks.append(_table_to_markdown(rows))
        return [TextPage(text="\n\n".join(blocks), page_number=None)]


def _normalise_page(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    out: list[str] = []
    blank = False
    for line in lines:
        if line.strip():
            out.append(line)
            blank = False
        elif not blank:
            out.append("")
            blank = True
    return "\n".join(out).strip()


def _remove_repeated_page_margins(pages: list[TextPage]) -> list[TextPage]:
    """Remove only exact short headers/footers repeated on most pages."""
    if len(pages) < 3:
        return pages
    edges: list[tuple[str, str]] = []
    for page in pages:
        lines = [line.strip() for line in page.text.splitlines() if line.strip()]
        edges.append((lines[0] if lines else "", lines[-1] if lines else ""))
    threshold = max(3, (len(pages) * 2 + 2) // 3)
    repeated = set()
    for position in (0, 1):
        counts: dict[str, int] = {}
        for edge in edges:
            value = edge[position]
            if value and len(value) <= 120:
                counts[value] = counts.get(value, 0) + 1
        repeated.update(value for value, count in counts.items() if count >= threshold)
    if not repeated:
        return pages
    cleaned = []
    for page in pages:
        lines = page.text.splitlines()
        while lines and lines[0].strip() in repeated: lines.pop(0)
        while lines and lines[-1].strip() in repeated: lines.pop()
        cleaned.append(TextPage("\n".join(lines).strip(), page.page_number,
                                page.origin_type, page.confidence))
    return cleaned


def _table_to_markdown(table: Iterable[Iterable[object]]) -> str:
    rows = [["" if cell is None else str(cell).replace("\n", " ").strip() for cell in row] for row in table]
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    header = rows[0]
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)
